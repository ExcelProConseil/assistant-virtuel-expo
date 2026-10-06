"""Importe le planning Excel (un onglet par mois) dans la base SQLite des commandes."""
import sys, re, datetime as dt
import openpyxl
import os, shutil, tempfile, pathlib
from base import connecter, remplacer_lignes_excel, remplacer_salaries

COLS = dict(type="A", recep="B", ar="C", client="D", vk="E", cde="F", nomenc="G", qte="H",
            delai="I", rec_prod="J", designation="K", etape="L", depart="O",
            h_devis="P", ca="S")

def val(ws, col, r):
    v = ws[f"{col}{r}"].value
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, str):
        v = v.strip().replace("\xa0", " ")
        return v or None
    return v

def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

def couleur(ws, ref):
    c = ws[ref]
    rgb = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
    return rgb[-6:] if isinstance(rgb, str) and rgb[-6:] not in ("000000", "FFFFFF") else None

def lire(chemin):
    """Une ligne par article du planning. Les colonnes suivent l'Excel (A..U)."""
    wb = openpyxl.load_workbook(chemin, data_only=True)
    lignes = []
    for idx, ws in enumerate(wb):
        if ws["B1"].value != "DATE RECEP CDE":      # ignore l'onglet CP (planning des salariés)
            continue
        client = cde = type_ = None
        groupe = 0
        for r in range(3, ws.max_row + 1):
            vk, nomenc, qte = val(ws, "E", r), val(ws, "G", r), num(val(ws, "H", r))
            if not vk or not nomenc or qte is None:
                continue
            if val(ws, "D", r):                      # nouvelle commande (ligne principale)
                client, cde, groupe = val(ws, "D", r), val(ws, "F", r), groupe + 1
            if val(ws, "A", r):
                type_ = val(ws, "A", r)
            elif val(ws, "D", r):
                type_ = None
            delai = val(ws, "I", r)
            rec_prod = val(ws, "J", r)
            h_unit, ca_unit = num(val(ws, "P", r)), num(val(ws, "S", r))
            h_tot = num(val(ws, "R", r))
            if h_unit is not None and h_unit > 200:          # colonne décalée dans l'Excel : on ignore
                h_unit = h_tot = None
            if h_tot is None and h_unit is not None:
                h_tot = h_unit * qte
            ca_total = num(val(ws, "T", r))
            if ca_total is None and ca_unit is not None:
                ca_total = ca_unit * qte
            lignes.append(dict(
                vk=str(vk), mois=ws.title, ordre=idx * 10000 + r, groupe=f"{ws.title}-{groupe}",
                client=client, cde=str(cde) if cde else None, type=type_,
                nomenc=str(nomenc), designation=val(ws, "K", r), qte=qte,
                date_recep=val(ws, "B", r), date_ar=val(ws, "C", r),
                delai=delai if isinstance(delai, str) else None,
                rec_prod=str(rec_prod) if rec_prod is not None else None,
                etape=str(val(ws, "L", r)) if val(ws, "L", r) is not None else None,
                depart=val(ws, "O", r), h_unit=h_unit, h_reelles=num(val(ws, "Q", r)), h_tot=h_tot,
                ca_unit=ca_unit, ca_total=ca_total, couleur=couleur(ws, f"K{r}"),
                principale=1 if val(ws, "D", r) else 0,
                livree=1 if couleur(ws, f"K{r}") in VERT else 0))
    return lignes

# Lignes colorées en vert dans l'Excel = commande terminée
VERT = {"81D41A", "92D050"}

# ---- Onglet « CP 2026 » : salariés, absences, jours de fermeture ----
NOMS_MOIS = ["janvier", "fevrier", "mars", "avril", "mai", "juin", "juillet", "aout",
             "septembre", "octobre", "novembre", "decembre"]
ABSENCE = {"FFBF00", "FFC000"}       # orange : absence / jour non travaillé
FERIE = {"81D41A", "92D050"}         # vert : congé payé (ou fermeture)

def _sans_accents(t):
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn").lower().strip()

def lire_salaries(chemin, annee=2026):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    ws = next((w for w in wb if w["B1"].value != "DATE RECEP CDE" and "CP" in w.title.upper()), None)
    if ws is None:
        return [], []
    employes, absences, vus = [], [], {}
    r = 1
    while r <= ws.max_row:
        titre = _sans_accents(str(ws.cell(r, 1).value or ""))
        mois = next((i + 1 for i, m in enumerate(NOMS_MOIS) if titre.startswith(m)), None)
        if mois:
            r += 2                                   # saute la ligne des jours de la semaine
            while ws.cell(r, 1).value:
                nom = str(ws.cell(r, 1).value).strip()
                vus.setdefault(nom, len(vus))
                for j in range(1, 32):
                    try:
                        jour = dt.date(annee, mois, j)
                    except ValueError:
                        break
                    rgb = couleur(ws, ws.cell(r, 1 + j).coordinate)
                    if rgb in ABSENCE:
                        absences.append((jour.isoformat(), nom, "absence"))
                    elif rgb in FERIE:
                        absences.append((jour.isoformat(), nom, "cp"))
                r += 1
        r += 1
    return [(n, o) for n, o in vus.items()], absences

PLANNING = os.environ.get("JPV_PLANNING", r"C:/Users/cmore/OneDrive/Bureau/commandes clt/Systeme/Claude outputs/2026 PLANNING COMMANDE.xlsx")

def reimporter_si_modifie(con):
    """Relit le planning dès que JPV-AGENT (ou toi) l'a enregistré. Retourne True si relu."""
    if not os.path.exists(PLANNING):
        return False
    m = str(os.path.getmtime(PLANNING))
    r = con.execute("SELECT valeur FROM meta WHERE cle='planning_mtime'").fetchone()
    if r and r[0] == m:
        return False
    try:                                       # on lit une copie : marche même si Excel est ouvert
        tmp = pathlib.Path(tempfile.mkdtemp()) / "planning.xlsx"
        with open(PLANNING, "rb") as src, open(tmp, "wb") as dst:
            shutil.copyfileobj(src, dst)
    except PermissionError:                    # copie refusée (OneDrive, fichier verrouillé) : lecture directe
        tmp = PLANNING
    remplacer_lignes_excel(con, lire(tmp))
    remplacer_salaries(con, *lire_salaries(tmp))
    con.execute("INSERT OR REPLACE INTO meta VALUES ('planning_mtime', ?)", (m,))
    con.commit()
    return True

if __name__ == "__main__":
    chemin = sys.argv[1]
    lignes = lire(chemin)
    con = connecter()
    remplacer_lignes_excel(con, lignes)
    remplacer_salaries(con, *lire_salaries(chemin))
    print(f"{len(lignes)} lignes importées")
